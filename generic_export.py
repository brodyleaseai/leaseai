"""
LeaseAI - Generic Export Module
Export lease data to CSV and Excel formats with all extracted fields
Phase 1C: Generic CSV/Excel Exports
"""

import csv
import io
from datetime import datetime
from typing import Dict, List, Optional
import zipfile


class GenericCSVExporter:
    """
    Export lease data to generic CSV format with all extracted fields
    Includes all 84+ fields with confidence scores
    """
    
    @staticmethod
    def export_lease(lease_data: Dict, extracted_fields: Dict) -> bytes:
        """
        Export a single lease to generic CSV format
        
        Args:
            lease_data: Dictionary containing lease metadata from 'leases' table
            extracted_fields: Dictionary of extracted field values
            
        Returns:
            CSV file as bytes
        """
        output = io.StringIO()
        writer = csv.writer(output)
        
        # Header row with metadata
        writer.writerow(['LeaseAI - Lease Abstract Export'])
        writer.writerow(['Generated:', datetime.now().strftime('%Y-%m-%d %I:%M %p')])
        writer.writerow(['Lease ID:', lease_data.get('id', '')])
        writer.writerow(['Property:', lease_data.get('property_address', 'N/A')])
        writer.writerow(['Tenant:', lease_data.get('tenant_name', 'N/A')])
        writer.writerow([''])  # Blank row
        
        # Column headers
        writer.writerow(['Field Name', 'Value', 'Confidence', 'Data Type'])
        
        # Sort fields alphabetically for consistency
        sorted_fields = sorted(extracted_fields.items(), key=lambda x: x[0])
        
        for field_name, field_data in sorted_fields:
            if isinstance(field_data, dict):
                value = field_data.get('value', 'NOT FOUND')
                confidence = field_data.get('confidence', 'UNKNOWN')
            else:
                value = field_data if field_data else 'NOT FOUND'
                confidence = 'N/A'
            
            # Determine data type
            data_type = GenericCSVExporter._determine_data_type(field_name, value)
            
            writer.writerow([field_name, value, confidence, data_type])
        
        # Get CSV content as bytes
        csv_content = output.getvalue()
        output.close()
        
        return csv_content.encode('utf-8')
    
    @staticmethod
    def _determine_data_type(field_name: str, value: str) -> str:
        """Determine the data type of a field based on name and value"""
        field_lower = field_name.lower()
        value_str = str(value).lower()
        
        if 'date' in field_lower:
            return 'Date'
        elif any(x in field_lower for x in ['rent', 'amount', 'cost', 'price', 'deposit', 'allowance', 'charge']):
            return 'Currency'
        elif any(x in field_lower for x in ['square', 'footage', 'area', 'size']):
            return 'Number (Area)'
        elif any(x in field_lower for x in ['rate', 'percent', '%']):
            return 'Percentage'
        elif any(x in field_lower for x in ['days', 'months', 'years', 'term']):
            return 'Duration'
        else:
            return 'Text'


class GenericExcelExporter:
    """
    Export lease data to formatted Excel (.xlsx) format
    Creates multi-sheet workbook with summary, fields, ASC 842, and critical dates
    """
    
    @staticmethod
    def export_lease(lease_data: Dict, extracted_fields: Dict, asc842_data: Optional[Dict] = None) -> bytes:
        """
        Export a single lease to Excel format with multiple sheets
        
        Args:
            lease_data: Dictionary containing lease metadata from 'leases' table
            extracted_fields: Dictionary of extracted field values
            asc842_data: Optional ASC 842 calculation results
            
        Returns:
            Excel file as bytes
        """
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        except ImportError:
            raise ImportError("openpyxl is required for Excel export. Install with: pip install openpyxl")
        
        wb = Workbook()
        
        # Remove default sheet
        wb.remove(wb.active)
        
        # Sheet 1: Lease Summary
        GenericExcelExporter._create_summary_sheet(wb, lease_data, extracted_fields)
        
        # Sheet 2: All Extracted Fields
        GenericExcelExporter._create_fields_sheet(wb, extracted_fields)
        
        # Sheet 3: ASC 842 Calculations (if available)
        if asc842_data:
            GenericExcelExporter._create_asc842_sheet(wb, asc842_data)
        
        # Sheet 4: Critical Dates
        GenericExcelExporter._create_critical_dates_sheet(wb, extracted_fields)
        
        # Save to bytes
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        
        return output.getvalue()
    
    @staticmethod
    def _create_summary_sheet(wb: 'Workbook', lease_data: Dict, extracted_fields: Dict):
        """Create Sheet 1: Lease Summary with key information"""
        from openpyxl.styles import Font, PatternFill, Alignment
        
        ws = wb.create_sheet("Lease Summary", 0)
        ws.column_dimensions['A'].width = 30
        ws.column_dimensions['B'].width = 50
        
        # Title
        ws['A1'] = 'LeaseAI - Lease Abstract'
        ws['A1'].font = Font(size=18, bold=True, color='1a1f2e')
        ws['A1'].fill = PatternFill(start_color='64b5f6', end_color='64b5f6', fill_type='solid')
        ws.merge_cells('A1:B1')
        
        # Generated date
        ws['A2'] = 'Generated:'
        ws['B2'] = datetime.now().strftime('%B %d, %Y at %I:%M %p')
        ws['A2'].font = Font(bold=True)
        
        # Key information
        row = 4
        summary_fields = [
            ('Lease ID', lease_data.get('id', 'N/A')),
            ('Property Address', lease_data.get('property_address', 'N/A')),
            ('Tenant Name', lease_data.get('tenant_name', 'N/A')),
            ('Status', lease_data.get('status', 'N/A')),
            ('', ''),  # Blank row
            ('Landlord Name', GenericExcelExporter._get_field_value(extracted_fields, 'Landlord Legal Name')),
            ('Tenant Legal Name', GenericExcelExporter._get_field_value(extracted_fields, 'Tenant Legal Name')),
            ('', ''),
            ('Lease Commencement Date', GenericExcelExporter._get_field_value(extracted_fields, 'Lease Commencement Date')),
            ('Lease Expiration Date', GenericExcelExporter._get_field_value(extracted_fields, 'Lease Expiration Date')),
            ('Base Term', GenericExcelExporter._get_field_value(extracted_fields, 'Base Lease Term (months or years)')),
            ('', ''),
            ('Base Rent (Year 1)', GenericExcelExporter._get_field_value(extracted_fields, 'Base Rent Year 1 (monthly)')),
            ('Square Footage', GenericExcelExporter._get_field_value(extracted_fields, 'Premises Square Footage (RSF)')),
            ('Security Deposit', GenericExcelExporter._get_field_value(extracted_fields, 'Security Deposit Amount')),
            ('', ''),
            ('Renewal Options', GenericExcelExporter._get_field_value(extracted_fields, 'Number of Renewal Options')),
            ('Early Termination Right', GenericExcelExporter._get_field_value(extracted_fields, 'Early Termination Right (Yes/No)')),
        ]
        
        for label, value in summary_fields:
            if label:  # Skip blank rows
                ws[f'A{row}'] = label
                ws[f'B{row}'] = value
                ws[f'A{row}'].font = Font(bold=True)
            row += 1
    
    @staticmethod
    def _create_fields_sheet(wb: 'Workbook', extracted_fields: Dict):
        """Create Sheet 2: All Extracted Fields with confidence scores"""
        from openpyxl.styles import Font, PatternFill, Alignment
        
        ws = wb.create_sheet("Extracted Fields")
        ws.column_dimensions['A'].width = 50
        ws.column_dimensions['B'].width = 40
        ws.column_dimensions['C'].width = 15
        ws.column_dimensions['D'].width = 15
        
        # Headers
        headers = ['Field Name', 'Value', 'Confidence', 'Data Type']
        for col, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill(start_color='2d3561', end_color='2d3561', fill_type='solid')
            cell.alignment = Alignment(horizontal='center')
        
        # Freeze header row
        ws.freeze_panes = 'A2'
        
        # Sort fields alphabetically
        sorted_fields = sorted(extracted_fields.items(), key=lambda x: x[0])
        
        row = 2
        for field_name, field_data in sorted_fields:
            if isinstance(field_data, dict):
                value = field_data.get('value', 'NOT FOUND')
                confidence = field_data.get('confidence', 'UNKNOWN')
            else:
                value = field_data if field_data else 'NOT FOUND'
                confidence = 'N/A'
            
            data_type = GenericCSVExporter._determine_data_type(field_name, value)
            
            ws[f'A{row}'] = field_name
            ws[f'B{row}'] = value
            ws[f'C{row}'] = confidence
            ws[f'D{row}'] = data_type
            
            # Color code confidence
            if confidence == 'HIGH':
                ws[f'C{row}'].fill = PatternFill(start_color='d4edda', end_color='d4edda', fill_type='solid')
            elif confidence == 'MEDIUM':
                ws[f'C{row}'].fill = PatternFill(start_color='fff3cd', end_color='fff3cd', fill_type='solid')
            elif confidence == 'LOW':
                ws[f'C{row}'].fill = PatternFill(start_color='f8d7da', end_color='f8d7da', fill_type='solid')
            
            row += 1
    
    @staticmethod
    def _create_asc842_sheet(wb: 'Workbook', asc842_data: Dict):
        """Create Sheet 3: ASC 842 Calculations"""
        from openpyxl.styles import Font, PatternFill, Alignment
        
        ws = wb.create_sheet("ASC 842 Calculations")
        ws.column_dimensions['A'].width = 35
        ws.column_dimensions['B'].width = 20
        
        # Title
        ws['A1'] = 'ASC 842 Lease Calculations'
        ws['A1'].font = Font(size=14, bold=True)
        ws.merge_cells('A1:B1')
        
        row = 3
        
        # Key metrics
        metrics = [
            ('Lease Classification', asc842_data.get('classification', 'N/A')),
            ('Present Value of Lease Payments', asc842_data.get('present_value', 'N/A')),
            ('Incremental Borrowing Rate', asc842_data.get('discount_rate', 'N/A')),
            ('ROU Asset (Initial)', asc842_data.get('rou_asset', 'N/A')),
            ('Lease Liability (Initial)', asc842_data.get('lease_liability', 'N/A')),
            ('', ''),
            ('Total Lease Payments', asc842_data.get('total_payments', 'N/A')),
            ('Lease Term (months)', asc842_data.get('lease_term_months', 'N/A')),
        ]
        
        for label, value in metrics:
            if label:
                ws[f'A{row}'] = label
                ws[f'B{row}'] = value
                ws[f'A{row}'].font = Font(bold=True)
            row += 1
        
        # Payment schedule (if available)
        if 'payment_schedule' in asc842_data:
            row += 2
            ws[f'A{row}'] = 'Payment Schedule'
            ws[f'A{row}'].font = Font(bold=True, size=12)
            row += 1
            
            # Headers
            schedule_headers = ['Period', 'Payment', 'Interest', 'Principal', 'Balance']
            for col, header in enumerate(schedule_headers, start=1):
                cell = ws.cell(row=row, column=col, value=header)
                cell.font = Font(bold=True)
            
            row += 1
            # Add schedule rows (placeholder - would need actual schedule data)
    
    @staticmethod
    def _create_critical_dates_sheet(wb: 'Workbook', extracted_fields: Dict):
        """Create Sheet 4: Critical Dates Timeline"""
        from openpyxl.styles import Font, PatternFill, Alignment
        
        ws = wb.create_sheet("Critical Dates")
        ws.column_dimensions['A'].width = 40
        ws.column_dimensions['B'].width = 20
        ws.column_dimensions['C'].width = 50
        
        # Headers
        headers = ['Event', 'Date', 'Notes']
        for col, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill(start_color='2d3561', end_color='2d3561', fill_type='solid')
        
        row = 2
        
        # Collect all date-related fields
        date_fields = [
            ('Lease Commencement Date', 'Lease Commencement Date', 'Lease begins'),
            ('Rent Commencement Date', 'Rent Commencement Date', 'Rent payments begin'),
            ('Lease Expiration Date', 'Lease Expiration Date', 'Lease ends'),
            ('Renewal Notice Deadline', 'Renewal Option Notice (days prior)', 'Notice required for renewal'),
            ('Early Termination Available', 'Early Termination Date', 'Early termination right (if applicable)'),
        ]
        
        for event_name, field_key, notes in date_fields:
            value = GenericExcelExporter._get_field_value(extracted_fields, field_key)
            if value and value != 'NOT FOUND':
                ws[f'A{row}'] = event_name
                ws[f'B{row}'] = value
                ws[f'C{row}'] = notes
                row += 1
    
    @staticmethod
    def _get_field_value(extracted_fields: Dict, field_name: str) -> str:
        """Helper to extract field value from dict or nested dict"""
        field_data = extracted_fields.get(field_name)
        if isinstance(field_data, dict):
            return field_data.get('value', 'NOT FOUND')
        return field_data if field_data else 'NOT FOUND'


class BulkGenericExporter:
    """
    Bulk export multiple leases to generic CSV or Excel formats
    Creates a ZIP file containing all lease exports
    """
    
    @staticmethod
    def export_all_leases_to_csv(leases: List[Dict], extracted_fields_list: List[Dict]) -> bytes:
        """
        Export all leases to individual CSV files, bundled in a ZIP
        
        Args:
            leases: List of lease metadata dictionaries
            extracted_fields_list: List of extracted fields for each lease
            
        Returns:
            ZIP file containing all CSV files as bytes
        """
        zip_buffer = io.BytesIO()
        
        with zipfile.ZipFile(zip_buffer, 'w', zipfile.ZIP_DEFLATED) as zip_file:
            for lease, extracted_fields in zip(leases, extracted_fields_list):
                # Generate filename
                property_name = lease.get('property_address', 'Unknown').replace(' ', '_')
                property_name = ''.join(c for c in property_name if c.isalnum() or c == '_')[:30]
                filename = f"LeaseAbstract_{property_name}_{lease.get('id', 'unknown')}.csv"
                
                # Export to CSV
                csv_data = GenericCSVExporter.export_lease(lease, extracted_fields)
                
                # Add to ZIP
                zip_file.writestr(filename, csv_data)
        
        zip_buffer.seek(0)
        return zip_buffer.getvalue()
    
    @staticmethod
    def export_all_leases_to_excel(leases: List[Dict], extracted_fields_list: List[Dict], 
                                    asc842_data_list: Optional[List[Dict]] = None) -> bytes:
        """
        Export all leases to individual Excel files, bundled in a ZIP
        
        Args:
            leases: List of lease metadata dictionaries
            extracted_fields_list: List of extracted fields for each lease
            asc842_data_list: Optional list of ASC 842 calculations for each lease
            
        Returns:
            ZIP file containing all Excel files as bytes
        """
        zip_buffer = io.BytesIO()
        
        with zipfile.ZipFile(zip_buffer, 'w', zipfile.ZIP_DEFLATED) as zip_file:
            for idx, (lease, extracted_fields) in enumerate(zip(leases, extracted_fields_list)):
                # Generate filename
                property_name = lease.get('property_address', 'Unknown').replace(' ', '_')
                property_name = ''.join(c for c in property_name if c.isalnum() or c == '_')[:30]
                filename = f"LeaseAbstract_{property_name}_{lease.get('id', 'unknown')}.xlsx"
                
                # Get ASC 842 data if available
                asc842_data = None
                if asc842_data_list and idx < len(asc842_data_list):
                    asc842_data = asc842_data_list[idx]
                
                # Export to Excel
                excel_data = GenericExcelExporter.export_lease(lease, extracted_fields, asc842_data)
                
                # Add to ZIP
                zip_file.writestr(filename, excel_data)
        
        zip_buffer.seek(0)
        return zip_buffer.getvalue()
    
    @staticmethod
    def export_all_leases_to_consolidated_csv(leases: List[Dict], extracted_fields_list: List[Dict]) -> bytes:
        """
        Export all leases to a single consolidated CSV file
        Each row represents one lease with key fields
        
        Args:
            leases: List of lease metadata dictionaries
            extracted_fields_list: List of extracted fields for each lease
            
        Returns:
            CSV file as bytes
        """
        output = io.StringIO()
        writer = csv.writer(output)
        
        # Define key fields to include (most important fields only)
        key_fields = [
            'Landlord Legal Name',
            'Tenant Legal Name',
            'Premises Street Address',
            'City',
            'State',
            'Lease Commencement Date',
            'Lease Expiration Date',
            'Base Lease Term (months or years)',
            'Base Rent Year 1 (monthly)',
            'Premises Square Footage (RSF)',
            'Security Deposit Amount',
            'Number of Renewal Options',
            'Renewal Option Term (months or years per option)',
            'Early Termination Right (Yes/No)',
        ]
        
        # Write header row
        headers = ['Lease ID', 'Property Address', 'Tenant Name', 'Status'] + key_fields
        writer.writerow(headers)
        
        # Write data rows
        for lease, extracted_fields in zip(leases, extracted_fields_list):
            row = [
                lease.get('id', ''),
                lease.get('property_address', ''),
                lease.get('tenant_name', ''),
                lease.get('status', ''),
            ]
            
            # Add key field values
            for field_name in key_fields:
                field_data = extracted_fields.get(field_name)
                if isinstance(field_data, dict):
                    value = field_data.get('value', 'NOT FOUND')
                else:
                    value = field_data if field_data else 'NOT FOUND'
                row.append(value)
            
            writer.writerow(row)
        
        # Get CSV content as bytes
        csv_content = output.getvalue()
        output.close()
        
        return csv_content.encode('utf-8')


# Convenience functions for easy import
def export_lease_to_csv(lease_data: Dict, extracted_fields: Dict) -> bytes:
    """Quick export single lease to CSV"""
    return GenericCSVExporter.export_lease(lease_data, extracted_fields)


def export_lease_to_excel(lease_data: Dict, extracted_fields: Dict, asc842_data: Optional[Dict] = None) -> bytes:
    """Quick export single lease to Excel"""
    return GenericExcelExporter.export_lease(lease_data, extracted_fields, asc842_data)


def export_all_leases_csv_zip(leases: List[Dict], extracted_fields_list: List[Dict]) -> bytes:
    """Quick bulk export to ZIP of CSVs"""
    return BulkGenericExporter.export_all_leases_to_csv(leases, extracted_fields_list)


def export_all_leases_excel_zip(leases: List[Dict], extracted_fields_list: List[Dict], 
                                 asc842_data_list: Optional[List[Dict]] = None) -> bytes:
    """Quick bulk export to ZIP of Excel files"""
    return BulkGenericExporter.export_all_leases_to_excel(leases, extracted_fields_list, asc842_data_list)


def export_all_leases_consolidated_csv(leases: List[Dict], extracted_fields_list: List[Dict]) -> bytes:
    """Quick bulk export to single consolidated CSV"""
    return BulkGenericExporter.export_all_leases_to_consolidated_csv(leases, extracted_fields_list)
